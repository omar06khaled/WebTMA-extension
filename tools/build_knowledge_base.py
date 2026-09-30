"""Build firstCallExamples_enriched.json (Layer 1 knowledge base) from the desk manual.

Usage:
    python tools/build_knowledge_base.py "Desk Manual 2027.xlsx" [--sheet "Trades and Task Codes"] [--out firstCallExamples_enriched.json]

Reads the task-code sheet (columns: Task Description, Task Code, DTPC, POLY, TEMPE, WEST,
RFMT-DTPC, RFMT-POLY, RFMT-TMPE, RFMT-WEST, Notes) and writes the same JSON shape the
backend already uses:

    { "<Category>": { "<Task Description>": { "taskCode", "trade", "custodial", "notes" } } }

Anything it could not map cleanly is printed as a warning at the end so a person can check it.
Requires: pip install openpyxl
"""

import argparse
import collections
import json
import re
import sys

import openpyxl

CAMPUSES = ["dtpc", "poly", "tempe", "west", "rfmtDtpc", "rfmtPoly", "rfmtTmpe", "rfmtWest"]

ZONE_OPTIONS = {
    "dtpc": ["DTPC-A01", "DTPC-A02"],
    "poly": ["POLY-A01"],
    "tempe": ["TMPE-A", "TMPE-B", "TMPE-C"],
    "west": ["WEST-A01"],
    "rfmtDtpc": ["RFMT-TRADE"],
    "rfmtPoly": ["RFMT-POLY"],
    "rfmtTmpe": ["RFMT-ZONE-1", "RFMT-ZONE-2", "RFMT-ZONE-3"],
    "rfmtWest": ["RFMT-LSCS"],
}
ZONE_INSTRUCTION = "Check Zone Guide for correct zone"
RFMT_ZONE_INSTRUCTION = "Check RFMT Zone Guide"

# Shorthand in the sheet -> trade value the knowledge base has always used.
TRADE_ALIASES = {
    "CONST": "CONST (Global)",
    "MVRS": "MVRS (Global)",
    "SIGN-GLOBAL": "SIGN (Global)",
    "TMPE-FIRE": "TMPE-FIRE (Global)",
    "EMS": "EMS Team (Global)",
    "MOM": "MOM (Global)",
    "ESTIMATE": "Estimate",
    "LOCK": "LOCK",
    "CUST": "CUST",
    "JANI": "JANI",
    "EHS": "EHS",
    "HVAC": "HVAC",
    "ELEV": "ELEV",
    "UDIST": "UDIST",
    "TMPE-PM": "TMPE-PM",
    "RFMT PC": "RFMT PC",
    "CENTRAL PLANT": "Central Plant",
    "SEE NOTE": "SEE NOTE",
    "WEST-CPLNT": "West-CPLNT",
    "POLY-GRND": "POLY-GRND",
    "TMPE-GRND": "TMPE-GRND",
    "WEST-GRND": "WEST-GRND",
    # New in the 2027 manual, used as written.
    "RFMT-ELECTRICAL": "RFMT-Electrical",
    "RFMT-HVAC": "RFMT-HVAC",
    "RFMT-PLUMBING": "RFMT-Plumbing",
    "EMS/RFMT-HVAC": "EMS/RFMT-HVAC",
    "CAPSTONE": "CAPSTONE",
    "ZERO": "ZERO",
    "SSC": "SSC",
    "PAINT": "Paint",
}

# The sheet writes plain "GRND"; the knowledge base has always used the campus grounds trade.
# Campuses missing here have no known grounds trade and keep "GRND" (reported as a warning).
GRND_BY_CAMPUS = {
    "poly": "POLY-GRND",
    "rfmtPoly": "POLY-GRND",
    "tempe": "TMPE-GRND",
    "rfmtTmpe": "TMPE-GRND",
    "west": "WEST-GRND",
}

# Rows that look like tasks but are really section headers.
CATEGORY_ROWS = {"Lock Shop"}


def clean(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def zone(campus, instruction):
    return {"type": "ZONE", "options": ZONE_OPTIONS[campus], "instruction": instruction}


def map_trade(raw, campus, warnings, where):
    key = raw.upper()
    if key in ("", "N/A"):
        return None
    if key in ("ZONE", "ZONE/EOSS", "EMS/ZONE"):
        return zone(campus, ZONE_INSTRUCTION)
    if key in ("RFMT ZONE", "RFMT-ZONE"):
        return zone(campus, RFMT_ZONE_INSTRUCTION)
    if key == "GRND":
        if campus in GRND_BY_CAMPUS:
            return GRND_BY_CAMPUS[campus]
        warnings[f"'GRND' for {campus} has no known campus grounds trade; kept as 'GRND'"].append(where)
        return "GRND"
    if key in TRADE_ALIASES:
        return TRADE_ALIASES[key]
    warnings[f"unrecognized trade value {raw!r} for {campus}; kept as-is"].append(where)
    return raw


def parse_task_code(raw):
    if raw == "":
        return None
    try:
        number = float(raw)
    except ValueError:
        return raw  # e.g. "ISAAC"
    return int(number) if number.is_integer() else number


def build(path, sheet_name):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        sys.exit(f"Sheet {sheet_name!r} not found. Sheets: {workbook.sheetnames}")

    knowledge_base = {}
    warnings = collections.defaultdict(list)  # message -> rows it applies to
    general_notes = []
    category = None
    seen_codes = {}

    for row_number, row in enumerate(workbook[sheet_name].iter_rows(values_only=True), start=1):
        cells = [clean(value) for value in row] + [""] * 11
        description, code_raw = cells[0], cells[1]
        trades_raw, notes = cells[2:10], cells[10]

        if not any(cells[:11]) or description.lower() == "task description":
            continue
        if description.lower().startswith("general notes"):
            general_notes.append((category, description))
            continue
        if description in CATEGORY_ROWS or (description and not any(cells[1:11])):
            category = description
            knowledge_base.setdefault(category, {})
            continue
        if category is None:
            warnings["task appears before any category; skipped"].append(f"row {row_number} {description!r}")
            continue

        where = f"row {row_number} {description!r}"
        trade = {campus: map_trade(raw, campus, warnings, where) for campus, raw in zip(CAMPUSES, trades_raw)}
        entry = {
            "taskCode": parse_task_code(code_raw),
            "trade": trade,
            "custodial": any(value in ("CUST", "JANI") for value in trade.values()),
            "notes": notes or None,
        }

        existing = knowledge_base[category].get(description)
        if existing is not None:
            if existing != entry:
                warnings["duplicate description with different values in the same category; kept the first"].append(where)
            continue
        knowledge_base[category][description] = entry

        if isinstance(entry["taskCode"], int):
            seen_codes.setdefault(entry["taskCode"], []).append(f"{category} / {description}")

    knowledge_base = {name: tasks for name, tasks in knowledge_base.items() if tasks}

    for code, owners in seen_codes.items():
        if len(owners) > 1:
            warnings[f"task code {code} is used by more than one task"].extend(owners)

    return knowledge_base, warnings, general_notes


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("--sheet", default="Trades and Task Codes")
    parser.add_argument("--out", default="firstCallExamples_enriched.json")
    args = parser.parse_args()

    knowledge_base, warnings, general_notes = build(args.workbook, args.sheet)
    with open(args.out, "w", encoding="utf-8") as handle:
        json.dump(knowledge_base, handle, indent=2, ensure_ascii=False)
        handle.write("\n")

    total = sum(len(tasks) for tasks in knowledge_base.values())
    print(f"Wrote {args.out}: {len(knowledge_base)} categories, {total} tasks")
    print(f"{len(general_notes)} 'General Notes' rows are not in the knowledge base (they stay in the manual)")
    if warnings:
        print(f"\n{len(warnings)} things to review:")
        for message, places in warnings.items():
            shown = ", ".join(places[:4]) + (f", +{len(places) - 4} more" if len(places) > 4 else "")
            print(f"  - {message}: {shown}")


if __name__ == "__main__":
    main()
